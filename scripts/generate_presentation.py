"""
Generates a full PowerPoint deck showcasing the SalesIQ AI/ML internship
assignment: problem statement, objectives, features, tech stack,
architecture, ML approach, results, live screenshots, deliverables, and
limitations.

Usage:
    python scripts/generate_presentation.py
"""
from __future__ import annotations

from pptx import Presentation
from pptx.util import Inches, Pt, Emu
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.enum.shapes import MSO_SHAPE, MSO_CONNECTOR
from pptx.oxml.ns import qn
import copy

# ---------------------------------------------------------------- palette --
CREAM = RGBColor(0xFA, 0xF7, 0xF2)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)
BORDER = RGBColor(0xE7, 0xE0, 0xD1)
TEXT_PRIMARY = RGBColor(0x22, 0x1F, 0x1A)
TEXT_SECONDARY = RGBColor(0x6B, 0x62, 0x55)
TEXT_MUTED = RGBColor(0x9C, 0x93, 0x84)
BRAND = RGBColor(0x4F, 0x46, 0xE5)
BRAND_DARK = RGBColor(0x37, 0x2F, 0xA6)
TEAL = RGBColor(0x0D, 0x94, 0x88)
VIOLET = RGBColor(0x7C, 0x3A, 0xED)
AMBER = RGBColor(0xB4, 0x53, 0x09)
ROSE = RGBColor(0xE1, 0x1D, 0x48)
EMERALD = RGBColor(0x05, 0x96, 0x69)

FONT = "Segoe UI"

SCREENSHOTS = "docs/screenshots"

prs = Presentation()
prs.slide_width = Inches(13.333)
prs.slide_height = Inches(7.5)
SW, SH = prs.slide_width, prs.slide_height
BLANK = prs.slide_layouts[6]


def add_slide():
    slide = prs.slides.add_slide(BLANK)
    bg = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, SW, SH)
    bg.fill.solid()
    bg.fill.fore_color.rgb = CREAM
    bg.line.fill.background()
    bg.shadow.inherit = False
    bg._element.spPr.append(bg._element.spPr.makeelement(qn("a:effectLst"), {}))
    slide.shapes._spTree.remove(bg._element)
    slide.shapes._spTree.insert(2, bg._element)
    return slide


def set_no_shadow(shape):
    el = shape._element.spPr
    existing = el.find(qn("a:effectLst"))
    if existing is None:
        el.append(el.makeelement(qn("a:effectLst"), {}))


def add_rect(slide, x, y, w, h, fill=WHITE, line=BORDER, line_w=0.75, radius=None, shadow=False):
    shape_type = MSO_SHAPE.ROUNDED_RECTANGLE if radius else MSO_SHAPE.RECTANGLE
    shape = slide.shapes.add_shape(shape_type, x, y, w, h)
    if radius is not None:
        try:
            shape.adjustments[0] = radius
        except Exception:
            pass
    shape.fill.solid()
    shape.fill.fore_color.rgb = fill
    if line is None:
        shape.line.fill.background()
    else:
        shape.line.color.rgb = line
        shape.line.width = Pt(line_w)
    if not shadow:
        set_no_shadow(shape)
    return shape


def add_text(slide, x, y, w, h, text, size=14, color=TEXT_PRIMARY, bold=False,
             italic=False, align=PP_ALIGN.LEFT, anchor=MSO_ANCHOR.TOP, font=FONT,
             line_spacing=1.0, wrap=True):
    box = slide.shapes.add_textbox(x, y, w, h)
    tf = box.text_frame
    tf.word_wrap = wrap
    tf.vertical_anchor = anchor
    tf.margin_left = 0
    tf.margin_right = 0
    tf.margin_top = 0
    tf.margin_bottom = 0
    lines = text.split("\n")
    for i, line in enumerate(lines):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = align
        p.line_spacing = line_spacing
        r = p.add_run()
        r.text = line
        r.font.size = Pt(size)
        r.font.bold = bold
        r.font.italic = italic
        r.font.name = font
        r.font.color.rgb = color
    return box


def add_bullets(slide, x, y, w, h, items, size=13, color=TEXT_SECONDARY,
                 bullet_color=BRAND, gap=6, bold_lead=False, font=FONT, line_spacing=1.15):
    box = slide.shapes.add_textbox(x, y, w, h)
    tf = box.text_frame
    tf.word_wrap = True
    tf.margin_left = 0
    tf.margin_right = 0
    tf.margin_top = 0
    tf.margin_bottom = 0
    for i, item in enumerate(items):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.line_spacing = line_spacing
        p.space_after = Pt(gap)
        r1 = p.add_run()
        r1.text = "→  "
        r1.font.size = Pt(size)
        r1.font.bold = True
        r1.font.name = font
        r1.font.color.rgb = bullet_color
        if isinstance(item, tuple):
            lead, rest = item
            r2 = p.add_run()
            r2.text = lead
            r2.font.size = Pt(size)
            r2.font.bold = True
            r2.font.name = font
            r2.font.color.rgb = TEXT_PRIMARY
            r3 = p.add_run()
            r3.text = rest
            r3.font.size = Pt(size)
            r3.font.name = font
            r3.font.color.rgb = color
        else:
            r2 = p.add_run()
            r2.text = item
            r2.font.size = Pt(size)
            r2.font.name = font
            r2.font.color.rgb = color
    return box


def add_badge(slide, x, y, size=Inches(0.5)):
    badge = add_rect(slide, x, y, size, size, fill=BRAND, line=None, radius=0.3)
    tf = badge.text_frame
    tf.margin_left = 0
    tf.margin_right = 0
    tf.margin_top = 0
    tf.margin_bottom = 0
    tf.vertical_anchor = MSO_ANCHOR.MIDDLE
    p = tf.paragraphs[0]
    p.alignment = PP_ALIGN.CENTER
    r = p.add_run()
    r.text = "S"
    r.font.size = Pt(int(size / Pt(1) * 0.42))
    r.font.bold = True
    r.font.name = FONT
    r.font.color.rgb = WHITE
    return badge


def header(slide, kicker, title, page_no, total, subtitle=None):
    add_badge(slide, Inches(0.55), Inches(0.42), Inches(0.42))
    add_text(slide, Inches(1.1), Inches(0.38), Inches(4), Inches(0.3),
              "SalesIQ", size=13, bold=True, color=TEXT_PRIMARY)
    add_text(slide, Inches(1.1), Inches(0.63), Inches(4), Inches(0.25),
              "ANALYTICS & FORECASTING", size=8, bold=True, color=TEXT_MUTED)

    add_text(slide, Inches(0.55), Inches(1.05), Inches(11), Inches(0.3),
              kicker.upper(), size=11, bold=True, color=BRAND)
    add_text(slide, Inches(0.55), Inches(1.33), Inches(11.5), Inches(0.7),
              title, size=28, bold=True, color=TEXT_PRIMARY)
    y = Inches(2.05)
    if subtitle:
        add_text(slide, Inches(0.55), Inches(1.95), Inches(11.5), Inches(0.4),
                  subtitle, size=13, color=TEXT_SECONDARY)
        y = Inches(2.45)

    line_y = Inches(1.9) if not subtitle else Inches(2.35)
    line = slide.shapes.add_connector(MSO_CONNECTOR.STRAIGHT, Inches(0.55), line_y, Inches(12.78), line_y)
    line.line.color.rgb = BORDER
    line.line.width = Pt(1)

    add_text(slide, Inches(12.0), Inches(7.12), Inches(0.8), Inches(0.3),
              f"{page_no:02d} / {total}", size=9, color=TEXT_MUTED, align=PP_ALIGN.RIGHT)
    add_text(slide, Inches(0.55), Inches(7.12), Inches(4), Inches(0.3),
              "AI/ML Internship Assignment", size=9, color=TEXT_MUTED)
    return y


def stat_card(slide, x, y, w, h, label, value, accent=BRAND):
    add_rect(slide, x, y, w, h, fill=WHITE, line=BORDER, radius=0.08)
    bar = add_rect(slide, x, y, Inches(0.06), h, fill=accent, line=None)
    add_text(slide, x + Inches(0.25), y + Inches(0.18), w - Inches(0.4), Inches(0.3),
              label.upper(), size=9.5, bold=True, color=TEXT_MUTED)
    add_text(slide, x + Inches(0.25), y + Inches(0.5), w - Inches(0.4), Inches(0.6),
              value, size=20, bold=True, color=TEXT_PRIMARY)


def chip(slide, x, y, text, fill, text_color=WHITE, w=None, size=10.5):
    w = w or Inches(0.18 + 0.078 * len(text))
    shape = add_rect(slide, x, y, w, Inches(0.34), fill=fill, line=None, radius=0.5)
    tf = shape.text_frame
    tf.margin_left = Pt(2)
    tf.margin_right = Pt(2)
    tf.vertical_anchor = MSO_ANCHOR.MIDDLE
    p = tf.paragraphs[0]
    p.alignment = PP_ALIGN.CENTER
    r = p.add_run()
    r.text = text
    r.font.size = Pt(size)
    r.font.bold = True
    r.font.name = FONT
    r.font.color.rgb = text_color
    return shape, w


def styled_table(slide, x, y, w, h, headers, rows, col_widths=None, header_fill=None,
                  highlight_row=None, font_size=11.5):
    header_fill = header_fill or RGBColor(0x2E, 0x28, 0x87)
    n_rows = len(rows) + 1
    n_cols = len(headers)
    gtable = slide.shapes.add_table(n_rows, n_cols, x, y, w, h).table
    if col_widths:
        for i, cw in enumerate(col_widths):
            gtable.columns[i].width = cw
    for i, htext in enumerate(headers):
        cell = gtable.cell(0, i)
        cell.fill.solid()
        cell.fill.fore_color.rgb = header_fill
        cell.margin_left = Pt(8)
        cell.margin_top = Pt(4)
        cell.margin_bottom = Pt(4)
        tf = cell.text_frame
        tf.vertical_anchor = MSO_ANCHOR.MIDDLE
        p = tf.paragraphs[0]
        r = p.add_run()
        r.text = htext
        r.font.size = Pt(font_size)
        r.font.bold = True
        r.font.color.rgb = WHITE
        r.font.name = FONT
    for ridx, row in enumerate(rows, start=1):
        is_hl = highlight_row is not None and (ridx - 1) == highlight_row
        for cidx, val in enumerate(row):
            cell = gtable.cell(ridx, cidx)
            cell.fill.solid()
            cell.fill.fore_color.rgb = RGBColor(0xF4, 0xEF, 0xE4) if is_hl else WHITE
            cell.margin_left = Pt(8)
            cell.margin_top = Pt(3)
            cell.margin_bottom = Pt(3)
            tf = cell.text_frame
            tf.vertical_anchor = MSO_ANCHOR.MIDDLE
            p = tf.paragraphs[0]
            r = p.add_run()
            r.text = str(val)
            r.font.size = Pt(font_size)
            r.font.bold = is_hl and cidx == 0
            r.font.color.rgb = TEXT_PRIMARY if cidx == 0 else TEXT_SECONDARY
            r.font.name = FONT
    return gtable


TOTAL = 23

# ============================================================ 1. TITLE ====
slide = add_slide()
add_rect(slide, 0, 0, SW, Inches(7.5), fill=RGBColor(0x1B, 0x18, 0x40), line=None)
for i, (cx, cy, r, col) in enumerate([
    (Inches(11.3), Inches(0.2), Inches(3.2), RGBColor(0x2E, 0x28, 0x87)),
    (Inches(0.2), Inches(6.6), Inches(2.6), RGBColor(0x1E, 0x3A, 0x5F)),
]):
    circ = slide.shapes.add_shape(MSO_SHAPE.OVAL, cx, cy, r, r)
    circ.fill.solid()
    circ.fill.fore_color.rgb = col
    circ.line.fill.background()
    set_no_shadow(circ)

add_badge(slide, Inches(0.9), Inches(0.85), Inches(0.7))
add_text(slide, Inches(1.8), Inches(0.83), Inches(4), Inches(0.4), "SalesIQ", size=20, bold=True, color=WHITE)
add_text(slide, Inches(1.8), Inches(1.25), Inches(5), Inches(0.3), "ANALYTICS & FORECASTING", size=10, bold=True, color=RGBColor(0xA5, 0xB4, 0xFC))

add_text(slide, Inches(0.9), Inches(2.7), Inches(11.5), Inches(1.6),
          "Intelligent Sales Analytics\n& Forecasting Platform", size=42, bold=True, color=WHITE, line_spacing=1.05)
add_text(slide, Inches(0.9), Inches(4.35), Inches(10.5), Inches(0.6),
          "An end-to-end AI/ML system: raw sales data → validation & cleaning → EDA →\nfeature engineering → forecasting (Baseline · SARIMA · XGBoost) → REST API → dashboard",
          size=15, color=RGBColor(0xC7, 0xD2, 0xFE), line_spacing=1.3)

chips = ["Python 3.11+", "FastAPI", "Next.js 14", "SQLite / PostgreSQL", "Docker", "XGBoost", "SARIMA"]
cx = Inches(0.9)
for c in chips:
    shape, w = chip(slide, cx, Inches(5.3), c, fill=RGBColor(0x2E, 0x28, 0x87), size=10.5)
    cx += w + Inches(0.15)

add_text(slide, Inches(0.9), Inches(6.6), Inches(8), Inches(0.4),
          "AI/ML Internship Assignment  ·  Project Showcase", size=12, color=RGBColor(0x9C, 0xA3, 0xAF))
add_text(slide, Inches(0.9), Inches(6.95), Inches(8), Inches(0.35),
          "github.com/icodervivek/SalesIQ", size=11, color=RGBColor(0x8F, 0x96, 0xEE))

# ============================================================ 2. AGENDA ===
slide = add_slide()
header(slide, "Agenda", "What this deck covers", 2, TOTAL)
agenda_items = [
    ("01", "Problem Statement & Objectives", "Why SalesIQ, and what it needs to demonstrate"),
    ("02", "Core Features", "Upload, validation, EDA, forecasting, dashboard"),
    ("03", "Tech Stack & Architecture", "How the system is put together"),
    ("04", "ML Approach", "Baseline, SARIMA, XGBoost, features, validation"),
    ("05", "Results", "Measured model comparison, not illustrative numbers"),
    ("06", "Live Product Walkthrough", "Screenshots of the running dashboard"),
    ("07", "Testing, Deliverables & Limitations", "What's verified, what's shipped, what's next"),
]
y = Inches(2.7)
for i, (num, title, desc) in enumerate(agenda_items):
    col = i % 2
    row = i // 2
    x = Inches(0.55) + col * Inches(6.15)
    yy = y + row * Inches(1.28)
    add_rect(slide, x, yy, Inches(5.85), Inches(1.08), fill=WHITE, line=BORDER, radius=0.1)
    add_text(slide, x + Inches(0.25), yy + Inches(0.18), Inches(0.9), Inches(0.7), num, size=22, bold=True, color=BORDER)
    add_text(slide, x + Inches(1.05), yy + Inches(0.16), Inches(4.6), Inches(0.4), title, size=14, bold=True, color=TEXT_PRIMARY)
    add_text(slide, x + Inches(1.05), yy + Inches(0.58), Inches(4.6), Inches(0.4), desc, size=10.5, color=TEXT_SECONDARY)

# ====================================================== 3. PROBLEM STMT ===
slide = add_slide()
header(slide, "The Challenge", "Problem Statement", 3, TOTAL)
add_rect(slide, Inches(0.55), Inches(2.6), Inches(7.1), Inches(4.1), fill=WHITE, line=BORDER, radius=0.06)
add_text(slide, Inches(0.95), Inches(2.95), Inches(6.3), Inches(1.5),
          "Businesses sit on large volumes of historical sales data, but converting "
          "that data into reliable, forward-looking information is hard.",
          size=15, color=TEXT_PRIMARY, line_spacing=1.35)
add_text(slide, Inches(0.95), Inches(4.15), Inches(6.3), Inches(1.2),
          "Manual spreadsheets can show what already happened — they are far less "
          "effective at surfacing seasonality, trend shifts, or expected future demand.",
          size=13, color=TEXT_SECONDARY, line_spacing=1.35)
add_rect(slide, Inches(0.95), Inches(5.5), Inches(6.3), Inches(0.95), fill=RGBColor(0xEE, 0xF0, 0xFD), line=None, radius=0.1)
add_text(slide, Inches(1.2), Inches(5.68), Inches(5.8), Inches(0.6),
          "SalesIQ's goal: build an end-to-end forecasting & analytics system that "
          "turns raw transaction data into actionable forecasts and visual insight.",
          size=12.5, bold=True, color=BRAND_DARK, line_spacing=1.3)

x2 = Inches(7.95)
add_rect(slide, x2, Inches(2.6), Inches(4.85), Inches(4.1), fill=RGBColor(0x1B, 0x18, 0x40), line=None, radius=0.06)
add_text(slide, x2 + Inches(0.4), Inches(2.9), Inches(4), Inches(0.35), "Who feels this pain", size=13, bold=True, color=WHITE)
users = [
    ("Sales Manager", "Understand performance & expected future sales"),
    ("Business Analyst", "Explore trends, products, regions, scenarios"),
    ("Operations Team", "Use demand forecasts for planning"),
    ("Technical Team", "Run, evaluate, maintain the ML pipeline"),
]
yy = Inches(3.45)
for name, need in users:
    add_text(slide, x2 + Inches(0.4), yy, Inches(4.1), Inches(0.3), name, size=12.5, bold=True, color=RGBColor(0xA5, 0xB4, 0xFC))
    add_text(slide, x2 + Inches(0.4), yy + Inches(0.32), Inches(4.1), Inches(0.4), need, size=10.5, color=RGBColor(0xD1, 0xD5, 0xDB))
    yy += Inches(0.78)

# =================================================== 4. OBJECTIVES ========
slide = add_slide()
header(slide, "Scope", "Project Objectives", 4, TOTAL)
objectives = [
    ("Ingestion & Cleaning", "A clean, reusable pipeline for sales-data validation and preprocessing"),
    ("Exploratory Analysis", "Surface trends, seasonality, outliers, and product/region performance"),
    ("Feature Engineering", "Time-based and business features (lags, rolling stats, calendar, exogenous)"),
    ("2+ Forecasting Models", "Implement and fairly compare multiple forecasting approaches"),
    ("Time-Aware Evaluation", "Chronological validation with forecasting-appropriate metrics"),
    ("REST API", "Expose data, analytics, training and forecasting through endpoints"),
    ("Interactive Dashboard", "A real UI for uploading, exploring, and forecasting"),
    ("Documentation & Packaging", "Assumptions, limitations, setup — runnable by another developer"),
]
cols = 2
for i, (title, desc) in enumerate(objectives):
    col = i % cols
    row = i // cols
    x = Inches(0.55) + col * Inches(6.15)
    y = Inches(2.55) + row * Inches(1.12)
    add_rect(slide, x, y, Inches(0.42), Inches(0.42), fill=BRAND, line=None, radius=0.3)
    tf = slide.shapes[-1].text_frame
    tf.vertical_anchor = MSO_ANCHOR.MIDDLE
    p = tf.paragraphs[0]
    p.alignment = PP_ALIGN.CENTER
    r = p.add_run(); r.text = "✓"; r.font.size = Pt(16); r.font.bold = True; r.font.color.rgb = WHITE
    add_text(slide, x + Inches(0.6), y - Inches(0.02), Inches(5.4), Inches(0.35), title, size=13.5, bold=True, color=TEXT_PRIMARY)
    add_text(slide, x + Inches(0.6), y + Inches(0.34), Inches(5.4), Inches(0.65), desc, size=10.5, color=TEXT_SECONDARY, line_spacing=1.2)

# =================================================== 5. CORE FEATURES =====
slide = add_slide()
header(slide, "What It Does", "Core Features", 5, TOTAL)
feat_headers = ["Feature", "Expected Behavior"]
feat_rows = [
    ["Data Upload", "Upload CSV sales data, validate required columns"],
    ["Validation & Cleaning", "Detect + handle missing/invalid dates, duplicates, negatives"],
    ["EDA Dashboard", "Trends, growth, product & regional performance, distributions"],
    ["Feature Engineering", "Date, lag, rolling, growth-rate, exogenous features"],
    ["Forecasting", "Selectable horizon: 7 / 30 / 90 / custom days"],
    ["Model Comparison", "Baseline · SARIMA · XGBoost, consistent metrics"],
    ["Prediction Intervals", "Shaded forecast range, not just a point estimate"],
    ["Product/Region Forecasting", "Scope training & forecasting to a segment"],
    ["Anomaly Detection", "Rolling z-score flags unusual sales periods"],
    ["CSV Export", "Download forecast results"],
]
styled_table(slide, Inches(0.55), Inches(2.5), Inches(12.25), Inches(4.5), feat_headers, feat_rows,
             col_widths=[Inches(3.4), Inches(8.85)], font_size=12.5)

# =================================================== 6. DATASET SPEC ======
slide = add_slide()
header(slide, "Input", "Dataset Specification", 6, TOTAL)
ds_headers = ["Column", "Type", "Requirement"]
ds_rows = [
    ["date", "Date", "Required"],
    ["product_id / product_name", "String", "Required"],
    ["units_sold / revenue", "Numeric", "Required"],
    ["category / region", "String", "Recommended"],
    ["unit_price", "Numeric", "Recommended"],
    ["discount / marketing_spend", "Numeric", "Optional"],
]
styled_table(slide, Inches(0.55), Inches(2.5), Inches(7.1), Inches(3.5), ds_headers, ds_rows,
             col_widths=[Inches(3.2), Inches(1.7), Inches(2.2)], font_size=12)

x2 = Inches(7.95)
add_rect(slide, x2, Inches(2.5), Inches(4.85), Inches(4.15), fill=WHITE, line=BORDER, radius=0.08)
add_text(slide, x2 + Inches(0.35), Inches(2.75), Inches(4.2), Inches(0.3), "6 bundled sample datasets", size=13, bold=True, color=TEXT_PRIMARY)
samples = [
    ("sample_sales.csv", "29,330 rows · main dataset, 24 months"),
    ("quick_test_sample.csv", "734 rows · fast single-product test"),
    ("fmcg_grocery_sample.csv", "10,220 rows · different vertical"),
    ("minimal_history_sample.csv", "50 rows · training boundary case"),
    ("sparse_gaps_sample.csv", "342 rows · missing calendar days"),
    ("edge_case_sample.csv", "13 rows · every cleaning rule at once"),
]
yy = Inches(3.2)
for name, desc in samples:
    add_text(slide, x2 + Inches(0.35), yy, Inches(4.2), Inches(0.25), name, size=11, bold=True, color=BRAND_DARK)
    add_text(slide, x2 + Inches(0.35), yy + Inches(0.24), Inches(4.2), Inches(0.25), desc, size=9.5, color=TEXT_SECONDARY)
    yy += Inches(0.55)

# =================================================== 7. ARCHITECTURE ======
slide = add_slide()
header(slide, "System Design", "Architecture", 7, TOTAL)

boxes = [
    (Inches(0.55), Inches(2.6), Inches(3.7), Inches(1.1), "Next.js Frontend", "Dashboard · Upload · Filters · Forecast viz", BRAND),
    (Inches(4.65), Inches(2.6), Inches(3.7), Inches(1.1), "FastAPI Backend", "REST endpoints · validation · orchestration", VIOLET),
    (Inches(8.75), Inches(2.6), Inches(3.7), Inches(1.1), "ML Pipeline", "preprocessing → features → models → eval", TEAL),
    (Inches(0.55), Inches(4.3), Inches(5.85), Inches(1.1), "Database (SQLite / PostgreSQL)", "datasets · training runs · forecast results", AMBER),
    (Inches(6.65), Inches(4.3), Inches(5.85), Inches(1.1), "Docker + Docker Compose", "reproducible dev & deployment runtime", EMERALD),
]
for x, y, w, h, title, sub, col in boxes:
    add_rect(slide, x, y, w, h, fill=WHITE, line=BORDER, radius=0.08)
    add_rect(slide, x, y, Inches(0.08), h, fill=col, line=None)
    add_text(slide, x + Inches(0.3), y + Inches(0.2), w - Inches(0.5), Inches(0.35), title, size=13.5, bold=True, color=TEXT_PRIMARY)
    add_text(slide, x + Inches(0.3), y + Inches(0.58), w - Inches(0.5), Inches(0.45), sub, size=10, color=TEXT_SECONDARY)

arrow_y = Inches(3.72)
for x in [Inches(2.35), Inches(6.45)]:
    arr = slide.shapes.add_shape(MSO_SHAPE.DOWN_ARROW, x, arrow_y, Inches(0.35), Inches(0.35))
    arr.fill.solid(); arr.fill.fore_color.rgb = TEXT_MUTED; arr.line.fill.background(); set_no_shadow(arr)

add_text(slide, Inches(0.55), Inches(5.75), Inches(12.2), Inches(0.9),
          "Frontend talks to the backend only via REST. The backend is a thin API layer over a service "
          "layer that calls a framework-agnostic ML pipeline (pure pandas / scikit-learn / statsmodels / "
          "xgboost) — independently testable, reusable from a notebook or batch job.",
          size=12, color=TEXT_SECONDARY, line_spacing=1.35)

# =================================================== 8. TECH STACK ========
slide = add_slide()
header(slide, "Built With", "Tech Stack", 8, TOTAL)
stack = [
    ("Language", "Python 3.11+", TEAL),
    ("Data", "Pandas · NumPy", BRAND),
    ("ML", "Scikit-learn · XGBoost", VIOLET),
    ("Time Series", "Statsmodels (SARIMA)", AMBER),
    ("Backend", "FastAPI · SQLAlchemy", BRAND),
    ("Database", "SQLite (local) · PostgreSQL (Docker)", ROSE),
    ("Frontend", "Next.js 14 · React 18 · TypeScript", TEAL),
    ("Charts", "Recharts", VIOLET),
    ("Styling", "Tailwind CSS", EMERALD),
    ("Containerization", "Docker · Docker Compose", BRAND),
    ("Testing", "Pytest (22 tests)", AMBER),
    ("Version Control", "Git + GitHub", TEXT_PRIMARY),
]
cols = 4
cw, ch = Inches(2.95), Inches(1.28)
gx, gy = Inches(0.55), Inches(2.5)
for i, (label, val, col) in enumerate(stack):
    r, c = divmod(i, cols)
    x = gx + c * (cw + Inches(0.12))
    y = gy + r * (ch + Inches(0.12))
    add_rect(slide, x, y, cw, ch, fill=WHITE, line=BORDER, radius=0.1)
    dot = slide.shapes.add_shape(MSO_SHAPE.OVAL, x + Inches(0.22), y + Inches(0.22), Inches(0.16), Inches(0.16))
    dot.fill.solid(); dot.fill.fore_color.rgb = col; dot.line.fill.background(); set_no_shadow(dot)
    add_text(slide, x + Inches(0.22), y + Inches(0.5), cw - Inches(0.4), Inches(0.25), label.upper(), size=9, bold=True, color=TEXT_MUTED)
    add_text(slide, x + Inches(0.22), y + Inches(0.76), cw - Inches(0.4), Inches(0.45), val, size=11.5, bold=True, color=TEXT_PRIMARY, line_spacing=1.1)

# =================================================== 9. ML APPROACH =======
slide = add_slide()
header(slide, "Forecasting", "ML Approach — Three Models", 9, TOTAL)
models = [
    ("Baseline", "Seasonal-Naive", "Forecasts day t as the value 7 days earlier. Zero training cost — the bar every other model must clear.", TEXT_MUTED),
    ("Model 1", "SARIMA(1,1,1)(1,1,1,7)", "Classical statistical model with an explicit weekly seasonal term. Native confidence intervals.", VIOLET),
    ("Model 2", "XGBoost + Exogenous Features", "Gradient-boosted trees over lag/rolling/calendar features plus discount, marketing spend, and a lagged category/region peer benchmark.", TEAL),
]
x = Inches(0.55)
w = Inches(4.03)
for tag, name, desc, col in models:
    add_rect(slide, x, Inches(2.55), w, Inches(3.9), fill=WHITE, line=BORDER, radius=0.08)
    add_rect(slide, x, Inches(2.55), w, Inches(0.55), fill=col, line=None, radius=0.08)
    add_text(slide, x + Inches(0.25), Inches(2.68), w - Inches(0.5), Inches(0.3), tag.upper(), size=10.5, bold=True, color=WHITE)
    add_text(slide, x + Inches(0.25), Inches(3.25), w - Inches(0.5), Inches(0.8), name, size=15.5, bold=True, color=TEXT_PRIMARY, line_spacing=1.1)
    add_text(slide, x + Inches(0.25), Inches(4.05), w - Inches(0.5), Inches(2.2), desc, size=11.5, color=TEXT_SECONDARY, line_spacing=1.35)
    x += w + Inches(0.15)

add_text(slide, Inches(0.55), Inches(6.7), Inches(12.2), Inches(0.4),
          "LSTM/GRU intentionally not used — higher overfitting risk on ~24 months of data and far less interpretable than XGBoost's feature importances.",
          size=10.5, italic=True, color=TEXT_MUTED)

# =================================================== 10. FEATURE ENG ======
slide = add_slide()
header(slide, "Inputs", "Feature Engineering", 10, TOTAL)
fe_groups = [
    ("Calendar", "day · day-of-week · week · month · quarter · year · is_weekend · is_month_start/end", BRAND),
    ("Lag Features", "lag-1, lag-7, lag-14, lag-30 (strictly past values only — no leakage)", VIOLET),
    ("Rolling Stats", "7 / 14 / 30-day rolling mean & standard deviation, shifted by 1 day", TEAL),
    ("Growth Rate", "Recent 7-day growth computed from lagged values only", AMBER),
    ("Exogenous", "discount, marketing_spend (planned/known inputs) + a one-day-lagged product/category/region peer-average benchmark", ROSE),
]
y = Inches(2.55)
for title, desc, col in fe_groups:
    add_rect(slide, Inches(0.55), y, Inches(0.14), Inches(0.72), fill=col, line=None)
    add_text(slide, Inches(0.9), y - Inches(0.02), Inches(3), Inches(0.35), title, size=13.5, bold=True, color=TEXT_PRIMARY)
    add_text(slide, Inches(0.9), y + Inches(0.32), Inches(11), Inches(0.4), desc, size=11.5, color=TEXT_SECONDARY)
    y += Inches(0.92)

add_rect(slide, Inches(0.55), Inches(6.5), Inches(12.25), Inches(0.65), fill=RGBColor(0xFD, 0xF3, 0xE8), line=None, radius=0.15)
add_text(slide, Inches(0.85), Inches(6.66), Inches(11.7), Inches(0.4),
          "⚠  Bug caught during development: an early version used the SAME-DAY peer average, which is a near-perfect proxy "
          "for the target itself. Lagging it by one day fixed the leakage.",
          size=10.5, bold=True, color=AMBER)

# =================================================== 11. VALIDATION =======
slide = add_slide()
header(slide, "Rigor", "Validation & Evaluation", 11, TOTAL)
add_rect(slide, Inches(0.55), Inches(2.5), Inches(5.9), Inches(4.3), fill=WHITE, line=BORDER, radius=0.06)
add_text(slide, Inches(0.9), Inches(2.75), Inches(5.2), Inches(0.35), "Chronological Holdout", size=15, bold=True, color=TEXT_PRIMARY)
add_bullets(slide, Inches(0.9), Inches(3.25), Inches(5.2), Inches(3.3), [
    "Last N days (default 14) held out as test set",
    "Models fit ONLY on data before that cutoff",
    "No shuffling, ever — the #1 leakage source in forecasting",
    "Test size auto-capped at 20% of series length",
    "Minimum 45 days of history required to train",
], size=12.5, gap=14)

x2 = Inches(6.85)
add_rect(slide, x2, Inches(2.5), Inches(5.95), Inches(4.3), fill=WHITE, line=BORDER, radius=0.06)
add_text(slide, x2 + Inches(0.35), Inches(2.75), Inches(5.2), Inches(0.35), "Metrics Reported", size=15, bold=True, color=TEXT_PRIMARY)
add_bullets(slide, x2 + Inches(0.35), Inches(3.25), Inches(5.3), Inches(3.3), [
    ("MAE — ", "average absolute error, in units"),
    ("RMSE — ", "penalizes large misses more heavily"),
    ("MAPE — ", "percentage error, 1-unit floor prevents zero-actual blowups"),
    ("R² — ", "reported for context, not used to pick the winner"),
    ("Selection rule — ", "lowest holdout RMSE wins"),
], size=12.5, gap=14)

# =================================================== 12. RESULTS ==========
slide = add_slide()
header(slide, "Real Measured Results", "Model Comparison", 12, TOTAL, subtitle="Global daily series · sample_sales.csv · 731 days · 14-day holdout")
res_headers = ["Model", "MAE", "RMSE", "MAPE", "R²"]
res_rows = [
    ["Baseline (Seasonal Naive)", "32.93", "41.21", "4.99%", "0.834"],
    ["SARIMA(1,1,1)(1,1,1,7)", "41.63", "48.68", "6.05%", "0.768"],
    ["XGBoost (Lag+Calendar+Exogenous)", "33.72", "36.97", "5.05%", "0.866"],
]
styled_table(slide, Inches(0.55), Inches(2.9), Inches(12.25), Inches(1.9), res_headers, res_rows,
             col_widths=[Inches(4.6), Inches(1.9), Inches(1.9), Inches(1.9), Inches(1.95)],
             highlight_row=2, font_size=13)
add_text(slide, Inches(0.9), Inches(5.05), Inches(0.5), Inches(0.4), "🏆", size=18)
add_text(slide, Inches(0.55), Inches(5.15), Inches(12.2), Inches(1.6),
          "XGBoost wins on RMSE (fewer large misses around demand spikes). Seasonal-Naive is unusually "
          "strong here (clean weekly seasonality in the data) — a healthy reminder that a fancier model "
          "isn't automatically much better than a good baseline. For a single product's series (smaller, "
          "noisier), SARIMA sometimes wins instead — model selection runs fresh on every training run "
          "rather than hard-coding a winner.",
          size=12.5, color=TEXT_SECONDARY, line_spacing=1.4)

# =================================================== 13. API ==============
slide = add_slide()
header(slide, "Interface", "REST API", 13, TOTAL)
api_headers = ["Method", "Endpoint", "Purpose"]
api_rows = [
    ["POST", "/api/datasets/upload", "Upload + validate + clean a CSV"],
    ["GET", "/api/analytics/summary", "KPIs — revenue, units, AOV, growth"],
    ["GET", "/api/analytics/trends", "Historical trend series"],
    ["POST", "/api/models/train", "Train & evaluate all 3 models"],
    ["POST", "/api/forecast", "Generate a forecast for a horizon"],
    ["GET", "/api/forecast/{id}/export", "Download forecast as CSV"],
    ["GET", "/api/health", "Service health check"],
]
styled_table(slide, Inches(0.55), Inches(2.5), Inches(12.25), Inches(4.3), api_headers, api_rows,
             col_widths=[Inches(1.6), Inches(4.15), Inches(6.5)], font_size=12.5)

# ============================================= 14-18. SCREENSHOTS =========
def screenshot_slide(page_no, kicker, title, img, subtitle=None, big=True):
    slide = add_slide()
    header(slide, kicker, title, page_no, TOTAL, subtitle=subtitle)
    top = Inches(2.55) if subtitle else Inches(2.15)
    frame_w, frame_h = Inches(12.25), SH - top - Inches(0.55)
    frame_x = Inches(0.55)
    add_rect(slide, frame_x, top, frame_w, frame_h, fill=WHITE, line=BORDER, radius=0.03)
    pad = Inches(0.12)
    slide.shapes.add_picture(img, frame_x + pad, top + pad, width=frame_w - 2 * pad, height=frame_h - 2 * pad)
    return slide

screenshot_slide(14, "Live Product", "Overview Dashboard",
                  f"{SCREENSHOTS}/01-overview.jpg",
                  subtitle="KPI cards, revenue trend, and top products — live data from the bundled sample dataset")

screenshot_slide(15, "Live Product", "Regional / Category Performance & Data Quality",
                  f"{SCREENSHOTS}/02-overview-regional-quality.png",
                  subtitle="Revenue breakdowns plus a transparent validation & cleaning report for every upload")

screenshot_slide(16, "Live Product", "Upload & Validation",
                  f"{SCREENSHOTS}/03-upload.png",
                  subtitle="Drag-and-drop CSV upload with the expected column reference")

screenshot_slide(17, "Live Product", "Model Comparison",
                  f"{SCREENSHOTS}/04-forecast-model-comparison.jpg",
                  subtitle="All 3 models evaluated on the same holdout — best model flagged automatically")

screenshot_slide(18, "Live Product", "Forecast Result",
                  f"{SCREENSHOTS}/05-forecast-result.png",
                  subtitle="Historical vs. forecast clearly distinguished, shaded prediction interval, CSV export")

# =================================================== 19. TESTING ==========
slide = add_slide()
header(slide, "Verification", "Testing, Edge Cases & Quality", 19, TOTAL)
stat_card(slide, Inches(0.55), Inches(2.5), Inches(2.95), Inches(1.3), "Backend Tests", "22 / 22 passing", TEAL)
stat_card(slide, Inches(3.65), Inches(2.5), Inches(2.95), Inches(1.3), "Sample Datasets", "6 datasets", BRAND)
stat_card(slide, Inches(6.75), Inches(2.5), Inches(2.95), Inches(1.3), "Edge Cases Handled", "10+", VIOLET)
stat_card(slide, Inches(9.85), Inches(2.5), Inches(2.95), Inches(1.3), "API Endpoints", "13", AMBER)

add_text(slide, Inches(0.55), Inches(4.15), Inches(6), Inches(0.35), "Covered by tests", size=13.5, bold=True, color=TEXT_PRIMARY)
add_bullets(slide, Inches(0.55), Inches(4.6), Inches(5.8), Inches(2.3), [
    "Validation & cleaning edge cases",
    "Feature-engineering leakage checks (lag_1 uses only the past)",
    "Zero-actual-safe MAPE and all evaluation metrics",
    "Full upload → train → forecast → export API flow",
    "Exogenous-feature persistence in recursive forecasting",
], size=11.5, gap=8)

add_text(slide, Inches(6.85), Inches(4.15), Inches(6), Inches(0.35), "Data edge cases handled", size=13.5, bold=True, color=TEXT_PRIMARY)
add_bullets(slide, Inches(6.85), Inches(4.6), Inches(5.9), Inches(2.3), [
    "Missing/invalid dates, duplicate rows",
    "Missing required fields, negative values",
    "Zero-sales days, extreme outliers (winsorized)",
    "Gaps in the calendar (missing days)",
    "Datasets right at the minimum-history boundary",
], size=11.5, gap=8)

# =================================================== 20. DELIVERABLES =====
slide = add_slide()
header(slide, "Shipped", "Deliverables Checklist", 20, TOTAL)
deliverables = [
    "Complete source-code repository (GitHub, public)",
    "Working frontend dashboard (Next.js)",
    "Working FastAPI backend",
    "ML training & forecasting pipeline",
    "6 sample datasets (main + edge cases)",
    "requirements.txt + Dockerfile x2 + docker-compose.yml",
    "README — architecture, setup, API reference, assumptions",
    "Model evaluation report with real comparison results",
    "Screenshots from a live running instance",
    "22 automated tests",
]
cols = 2
for i, item in enumerate(deliverables):
    col = i % cols
    row = i // cols
    x = Inches(0.55) + col * Inches(6.15)
    y = Inches(2.5) + row * Inches(0.75)
    check = slide.shapes.add_shape(MSO_SHAPE.OVAL, x, y, Inches(0.32), Inches(0.32))
    check.fill.solid(); check.fill.fore_color.rgb = EMERALD; check.line.fill.background(); set_no_shadow(check)
    tf = check.text_frame; tf.vertical_anchor = MSO_ANCHOR.MIDDLE
    p = tf.paragraphs[0]; p.alignment = PP_ALIGN.CENTER
    r = p.add_run(); r.text = "✓"; r.font.size = Pt(13); r.font.bold = True; r.font.color.rgb = WHITE
    add_text(slide, x + Inches(0.48), y + Inches(0.02), Inches(5.5), Inches(0.5), item, size=12, color=TEXT_PRIMARY)

# =================================================== 21. BONUS ============
slide = add_slide()
header(slide, "Beyond the Baseline", "Bonus Features Implemented", 21, TOTAL)
done = [
    "Product & region-level forecasting",
    "Automatic model selection (by holdout RMSE)",
    "Prediction intervals / uncertainty visualization",
    "Anomaly detection (rolling z-score)",
    "CSV export of forecasts",
    "Docker Compose (frontend + backend + Postgres)",
]
not_done = [
    "What-if analysis (discount/marketing scenarios)",
    "Scheduled retraining / background job queue",
    "Authentication & role-based access",
    "Cloud deployment",
    "MLflow experiment tracking",
]
add_rect(slide, Inches(0.55), Inches(2.5), Inches(6.1), Inches(4.3), fill=WHITE, line=BORDER, radius=0.06)
add_text(slide, Inches(0.85), Inches(2.7), Inches(5.5), Inches(0.35), "✓  Implemented", size=14, bold=True, color=EMERALD)
add_bullets(slide, Inches(0.85), Inches(3.15), Inches(5.5), Inches(3.5), done, size=12, bullet_color=EMERALD, gap=12)

x2 = Inches(6.85)
add_rect(slide, x2, Inches(2.5), Inches(5.95), Inches(4.3), fill=WHITE, line=BORDER, radius=0.06)
add_text(slide, x2 + Inches(0.35), Inches(2.7), Inches(5.4), Inches(0.35), "○  Explicitly out of scope", size=14, bold=True, color=TEXT_MUTED)
add_bullets(slide, x2 + Inches(0.35), Inches(3.15), Inches(5.4), Inches(3.5), not_done, size=12, bullet_color=TEXT_MUTED, color=TEXT_SECONDARY, gap=12)

# =================================================== 22. LIMITATIONS ======
slide = add_slide()
header(slide, "Honest Assessment", "Limitations & Production Roadmap", 22, TOTAL)
add_text(slide, Inches(0.55), Inches(2.5), Inches(6), Inches(0.35), "Known limitations", size=14, bold=True, color=ROSE)
add_bullets(slide, Inches(0.55), Inches(2.95), Inches(5.9), Inches(3.8), [
    "Single-fold holdout, not full walk-forward backtesting",
    "Missing calendar days treated as zero sales",
    "XGBoost's interval is a residual-based approximation",
    "No cold-start model for brand-new products",
    "No holiday calendar or competitor pricing signals",
], size=11.5, bullet_color=ROSE, gap=10)

x2 = Inches(6.85)
add_text(slide, x2, Inches(2.5), Inches(6), Inches(0.35), "Production improvement path", size=14, bold=True, color=EMERALD)
add_bullets(slide, x2, Inches(2.95), Inches(5.9), Inches(3.8), [
    "Expanding-window backtesting as a background job",
    "Holiday calendar + competitor pricing as features",
    "Per-segment models with fallback to the global model",
    "MLflow experiment tracking",
    "Drift monitoring — alert when live MAPE regresses",
], size=11.5, bullet_color=EMERALD, gap=10)

# =================================================== 23. THANK YOU ========
slide = add_slide()
add_rect(slide, 0, 0, SW, Inches(7.5), fill=RGBColor(0x1B, 0x18, 0x40), line=None)
add_badge(slide, Inches(5.87), Inches(2.15), Inches(0.8))
add_text(slide, Inches(0.9), Inches(3.2), Inches(11.5), Inches(0.7), "Thank You", size=40, bold=True, color=WHITE, align=PP_ALIGN.CENTER)
add_text(slide, Inches(0.9), Inches(4.0), Inches(11.5), Inches(0.5),
          "SalesIQ — from raw sales data to a validated forecast, end to end.",
          size=15, color=RGBColor(0xC7, 0xD2, 0xFE), align=PP_ALIGN.CENTER)
add_text(slide, Inches(0.9), Inches(4.9), Inches(11.5), Inches(0.4),
          "github.com/icodervivek/SalesIQ", size=14, bold=True, color=RGBColor(0xA5, 0xB4, 0xFC), align=PP_ALIGN.CENTER)
add_text(slide, Inches(0.9), Inches(5.4), Inches(11.5), Inches(0.4),
          "Full docs, model evaluation report, and setup instructions in the repo README",
          size=11.5, color=RGBColor(0x9C, 0xA3, 0xAF), align=PP_ALIGN.CENTER)

out_path = "SalesIQ_Assignment_Presentation.pptx"
prs.save(out_path)
print(f"Saved {out_path} with {len(prs.slides.__iter__.__self__._sldIdLst)} slides")
