"""
Generates a plain-language Word report explaining the SalesIQ project:
problem statement, features, tech stack, high-level design, data flow,
a screenshot-by-screenshot dashboard walkthrough, and the ML approach.

Usage:
    python scripts/generate_report_docx.py
"""
from __future__ import annotations

from docx import Document
from docx.shared import Pt, Inches, RGBColor, Cm
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

SCREENSHOTS = "docs/screenshots"

BRAND = RGBColor(0x4F, 0x46, 0xE5)
BRAND_DARK = RGBColor(0x37, 0x2F, 0xA6)
TEXT_PRIMARY = RGBColor(0x22, 0x1F, 0x1A)
TEXT_SECONDARY = RGBColor(0x4A, 0x44, 0x3A)
HEADER_FILL = "372FA6"
ROW_ALT_FILL = "F4EFE4"

doc = Document()

# ------------------------------------------------------------- base style --
style = doc.styles["Normal"]
style.font.name = "Calibri"
style.font.size = Pt(11)
style.font.color.rgb = TEXT_PRIMARY
style.paragraph_format.space_after = Pt(8)
style.paragraph_format.line_spacing = 1.3

for section in doc.sections:
    section.top_margin = Cm(2.0)
    section.bottom_margin = Cm(2.0)
    section.left_margin = Cm(2.2)
    section.right_margin = Cm(2.2)


def set_cell_shading(cell, hex_color):
    tcPr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), hex_color)
    tcPr.append(shd)


def h1(text):
    p = doc.add_heading(level=1)
    r = p.add_run(text)
    r.font.color.rgb = BRAND_DARK
    r.font.size = Pt(20)
    p.paragraph_format.space_before = Pt(20)
    p.paragraph_format.space_after = Pt(10)
    return p


def h2(text):
    p = doc.add_heading(level=2)
    r = p.add_run(text)
    r.font.color.rgb = BRAND
    r.font.size = Pt(15)
    p.paragraph_format.space_before = Pt(14)
    p.paragraph_format.space_after = Pt(6)
    return p


def body(text, size=11, italic=False, color=TEXT_PRIMARY, space_after=8):
    p = doc.add_paragraph()
    r = p.add_run(text)
    r.font.size = Pt(size)
    r.italic = italic
    r.font.color.rgb = color
    p.paragraph_format.space_after = Pt(space_after)
    return p


def bullet(text, bold_lead=None):
    p = doc.add_paragraph(style="List Bullet")
    if bold_lead:
        r1 = p.add_run(bold_lead)
        r1.bold = True
        r1.font.color.rgb = TEXT_PRIMARY
        r2 = p.add_run(text)
    else:
        r2 = p.add_run(text)
    r2.font.size = Pt(11)
    p.paragraph_format.space_after = Pt(4)
    return p


def numbered(text, bold_lead=None):
    p = doc.add_paragraph(style="List Number")
    if bold_lead:
        r1 = p.add_run(bold_lead)
        r1.bold = True
        r2 = p.add_run(text)
    else:
        r2 = p.add_run(text)
    r2.font.size = Pt(11)
    p.paragraph_format.space_after = Pt(6)
    return p


def table(headers, rows, widths=None):
    t = doc.add_table(rows=1, cols=len(headers))
    t.style = "Table Grid"
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    hdr = t.rows[0].cells
    for i, htext in enumerate(headers):
        hdr[i].text = ""
        p = hdr[i].paragraphs[0]
        r = p.add_run(htext)
        r.bold = True
        r.font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF)
        r.font.size = Pt(10.5)
        set_cell_shading(hdr[i], HEADER_FILL)
    for ridx, row in enumerate(rows):
        cells = t.add_row().cells
        for cidx, val in enumerate(row):
            cells[cidx].text = ""
            p = cells[cidx].paragraphs[0]
            r = p.add_run(str(val))
            r.font.size = Pt(10.5)
            if ridx % 2 == 1:
                set_cell_shading(cells[cidx], ROW_ALT_FILL)
    if widths:
        for i, w in enumerate(widths):
            for row in t.rows:
                row.cells[i].width = w
    doc.add_paragraph().paragraph_format.space_after = Pt(4)
    return t


def add_screenshot(path, caption, width=Inches(6.1)):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = p.add_run()
    run.add_picture(path, width=width)
    cap = doc.add_paragraph()
    cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = cap.add_run(caption)
    r.italic = True
    r.font.size = Pt(9.5)
    r.font.color.rgb = TEXT_SECONDARY
    cap.paragraph_format.space_after = Pt(14)


def page_break():
    doc.add_page_break()


# ==================================================================== COVER
title_p = doc.add_paragraph()
title_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
title_p.paragraph_format.space_before = Pt(100)
r = title_p.add_run("SalesIQ")
r.font.size = Pt(40)
r.bold = True
r.font.color.rgb = BRAND_DARK

sub_p = doc.add_paragraph()
sub_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
r = sub_p.add_run("Intelligent Sales Analytics & Forecasting Platform")
r.font.size = Pt(18)
r.font.color.rgb = TEXT_PRIMARY

sub2_p = doc.add_paragraph()
sub2_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
sub2_p.paragraph_format.space_before = Pt(10)
r = sub2_p.add_run("Project Report — AI/ML Internship Assignment")
r.font.size = Pt(13)
r.italic = True
r.font.color.rgb = TEXT_SECONDARY

sub3_p = doc.add_paragraph()
sub3_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
sub3_p.paragraph_format.space_before = Pt(60)
r = sub3_p.add_run("A plain-language walkthrough of what was built, why, and how it works.")
r.font.size = Pt(11)
r.font.color.rgb = TEXT_SECONDARY

link_p = doc.add_paragraph()
link_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
link_p.paragraph_format.space_before = Pt(200)
r = link_p.add_run("github.com/icodervivek/SalesIQ")
r.font.size = Pt(11)
r.font.color.rgb = BRAND

page_break()

# ============================================================ 1. PROBLEM ==
h1("1. The Problem, In Plain Terms")
body(
    "Most businesses have years of sales data sitting in spreadsheets or databases. That data can "
    "tell you what already happened — but turning it into a useful answer to “what will we sell "
    "next month?” is genuinely hard to do by hand. You'd need to notice seasonal patterns, spot "
    "which products are trending up or down, and build some kind of statistical model — none of which "
    "a spreadsheet does for you automatically."
)
body(
    "SalesIQ solves this by taking raw sales transactions, cleaning them up, and running them through "
    "a machine learning pipeline that produces an actual forecast — with a dashboard on top so anyone, "
    "not just a data scientist, can use it."
)
h2("Who is this for?")
table(
    ["Person", "What they need"],
    [
        ["Sales / Business Manager", "A quick read on performance and what's coming next"],
        ["Business Analyst", "The ability to slice data by product, region, or category"],
        ["Operations / Inventory Team", "Demand forecasts to plan stock and staffing"],
        ["Technical Team", "A pipeline they can run, evaluate, and maintain"],
    ],
    widths=[Inches(2.4), Inches(3.9)],
)

# ========================================================= 2. WHAT IS IT ==
h1("2. What SalesIQ Actually Does")
body(
    "Think of SalesIQ as three things stacked on top of each other, each doing one job well:"
)
numbered(" — someone drops a CSV of sales history into the app. SalesIQ checks it for problems "
         "(missing values, bad dates, duplicates) and cleans it up automatically.", "A data intake step ")
numbered(" — the cleaned data gets summarized into charts: revenue over time, top products, best "
         "regions, and so on. This is the part that answers “what already happened?”",
         "An analytics dashboard ")
numbered(" — three different forecasting models are trained on the historical data, compared "
         "side by side, and the best one is used to predict the next 7, 30, or 90 days. This answers "
         "“what's likely to happen next?”", "A forecasting engine ")
body(
    "All of this is wrapped in a normal web app — no command line or notebook needed to use it day to day."
)

h2("Feature List")
table(
    ["Feature", "What it does"],
    [
        ["Data Upload", "Upload a CSV; the app validates required columns are present"],
        ["Data Cleaning", "Fixes/removes bad rows: missing values, duplicates, negative numbers, bad dates"],
        ["EDA Dashboard", "Revenue trend, growth rate, top products, region/category breakdowns"],
        ["Feature Engineering", "Builds the inputs the ML models actually learn from (behind the scenes)"],
        ["Forecasting", "Pick a horizon (7 / 30 / 90 days, or a custom number) and get a forecast"],
        ["Model Comparison", "See all 3 models' accuracy side by side, with the best one highlighted"],
        ["Prediction Intervals", "A shaded “confidence range” around the forecast, not just one number"],
        ["Product/Region Scoping", "Forecast a specific product or region instead of the whole business"],
        ["Anomaly Detection", "Flags days where sales were unusually high or low"],
        ["CSV Export", "Download the forecast to use elsewhere (Excel, another tool, etc.)"],
    ],
    widths=[Inches(2.0), Inches(4.3)],
)

page_break()

# =========================================================== 3. TECH STACK
h1("3. Tech Stack — and Why Each Piece Was Chosen")
body("Every layer was picked for a specific, practical reason rather than “because it's popular”:")
table(
    ["Layer", "Technology", "Why"],
    [
        ["Backend API", "FastAPI (Python)", "Fast to build, auto-generates interactive API docs"],
        ["ML pipeline", "Pandas, scikit-learn, statsmodels, XGBoost", "The standard, well-tested tools for this kind of work"],
        ["Database", "SQLite by default, PostgreSQL via Docker", "Zero setup for local dev; production-grade option available"],
        ["Frontend", "Next.js + React + TypeScript", "A modern, fast, type-safe way to build the dashboard"],
        ["Charts", "Recharts", "Clean, composable charts that fit a React app naturally"],
        ["Styling", "Tailwind CSS", "Fast to build a consistent, professional-looking UI with"],
        ["Packaging", "Docker + Docker Compose", "So anyone can run the whole stack with one command"],
        ["Testing", "Pytest", "22 automated tests covering the pipeline and the API"],
    ],
    widths=[Inches(1.6), Inches(2.5), Inches(3.0)],
)

# ================================================================ 4. HLD ==
h1("4. High-Level Design (HLD)")
body(
    "The system is split into three independent layers that only talk to each other through clean "
    "interfaces — this means each one can be tested, changed, or even replaced without breaking the "
    "others."
)
numbered(" — the Next.js dashboard. It never touches the database directly; it only calls the "
         "backend's REST API, the same way any external client would.", "Frontend ")
numbered(" — a FastAPI application. Its only jobs are: validate incoming requests, call the "
         "right piece of ML code, save the result, and send back JSON.", "Backend ")
numbered(" — plain Python code (no web framework in it at all) that does the actual data "
         "cleaning, feature building, model training, and forecasting. Because it doesn't know "
         "anything about FastAPI or HTTP, it can be tested on its own or reused from a notebook.",
         "ML pipeline ")
body(
    "A small database (SQLite for local use, PostgreSQL when running with Docker) stores uploaded "
    "datasets, training results, and forecast history so they don't have to be recomputed every time."
)
h2("In one picture")
body(
    "Browser  →  Next.js Frontend  →  REST API (FastAPI)  →  ML Pipeline (pandas / scikit-learn / "
    "statsmodels / XGBoost)  →  Database\n"
    "                                                        ↕\n"
    "                                                Docker packages all of this together",
    size=10.5, italic=True, color=TEXT_SECONDARY
)

# =========================================================== 5. DATA FLOW
h1("5. How a Request Actually Flows Through the System")
body("Here's what happens, step by step, from uploading a file to seeing a forecast on screen:")
numbered(" — a CSV lands on the Upload page and is sent to the backend.", "Upload ")
numbered(" — the backend checks the required columns exist and reports back what's "
         "wrong, if anything (bad dates, missing fields, etc.).", "Validate ")
numbered(" — rows with unrecoverable problems are dropped; recoverable gaps (like a "
         "missing unit price) are filled in sensibly; extreme outliers are dampened rather than deleted.",
         "Clean ")
numbered(" — the cleaned dataset is saved, and a summary of what was found/fixed is shown back "
         "to the user immediately.", "Store & report ")
numbered(" — on the Overview page, the backend aggregates the clean data into KPIs and "
         "charts on request.", "Analyze ")
numbered(" — the backend builds calendar/lag/rolling features from the daily sales "
         "numbers, trains all three forecasting models on data before a cutoff date, and tests them "
         "against the days after that cutoff — never letting the models see the future during training.",
         "Train ")
numbered(" — using the best (or a chosen) model, the backend predicts forward for "
         "the requested horizon and returns both a forecast line and an uncertainty range.", "Forecast ")
numbered(" — the frontend draws the forecast chart, distinguishing historical data from "
         "the predicted future, and offers a CSV download.", "Display ")

page_break()

# ==================================================== 6. DASHBOARD WALKTHROUGH
h1("6. Dashboard Walkthrough")
body(
    "The screenshots below are from a live run of the app against the bundled sample dataset — "
    "not mockups."
)

h2("6.1 Overview Page — “What already happened”")
body(
    "This is the home page. Four KPI cards up top give the headline numbers — total revenue, units "
    "sold, average order value, and how many products are being tracked, along with a week-over-week "
    "growth indicator. Below that, a revenue trend chart shows the whole history at a glance, and a "
    "“Top Products” list ranks what's selling best."
)
add_screenshot(f"{SCREENSHOTS}/01-overview.jpg", "Overview — KPI cards and the revenue trend chart")

body(
    "Scrolling down, two bar charts break revenue down by region and by category, so you can spot "
    "which parts of the business are carrying the most weight. Underneath that is a Data Quality "
    "panel — this is the app being transparent about what it found and fixed in your uploaded file "
    "(missing values, duplicates, and so on), rather than silently changing your data."
)
add_screenshot(f"{SCREENSHOTS}/02-overview-regional-quality.png", "Regional/category performance and the data-quality report")

h2("6.2 Upload Page — Getting your data in")
body(
    "A simple drag-and-drop zone accepts a CSV file. Below it, a reference table lists exactly which "
    "columns are required, which are recommended, and which are optional, so there's no guessing about "
    "the expected format before you upload."
)
add_screenshot(f"{SCREENSHOTS}/03-upload.png", "Upload page with the expected CSV column reference")

h2("6.3 Forecasting Page — “What's likely to happen next”")
body(
    "This page has three steps. First, choose a scope (all data, or a specific product/region) and "
    "click Train Models — this runs the baseline, SARIMA, and XGBoost models and shows a comparison "
    "table with accuracy metrics (MAE, RMSE, MAPE, R²), with a trophy marking the best-performing "
    "model. A chart underneath shows how that model's predictions lined up against actual historical "
    "values it wasn't trained on."
)
add_screenshot(f"{SCREENSHOTS}/04-forecast-model-comparison.jpg", "All three models compared on the same holdout data")

body(
    "Second, pick a forecast horizon (7, 30, 90 days, or a custom number) and a model, then click "
    "Generate Forecast. The resulting chart clearly separates the solid historical line from the "
    "dashed future forecast, with a shaded band showing the range of likely outcomes — and a button "
    "to export the numbers as a CSV."
)
add_screenshot(f"{SCREENSHOTS}/05-forecast-result.png", "A 30-day forecast with historical data and a prediction interval")

page_break()

# ================================================== 7. ML APPROACH SIMPLE
h1("7. The Machine Learning, Explained Simply")
body(
    "Three different forecasting approaches are trained and compared every time, rather than "
    "committing to just one:"
)
h2("The Baseline — Seasonal Naive")
body(
    "The simplest possible forecast: “assume today looks like the same day last week.” It costs "
    "nothing to compute, but it's surprisingly hard to beat when sales have a strong weekly rhythm "
    "(which most retail data does). Its whole purpose is to be a fair bar that the smarter models have "
    "to clear — if a complex model can't beat this, it isn't actually adding value."
)
h2("Model 1 — SARIMA")
body(
    "A classical statistics model that explicitly models trend and seasonality (in this case, a "
    "7-day weekly cycle). It's the kind of model a statistician would reach for, and it comes with "
    "built-in, mathematically principled confidence intervals."
)
h2("Model 2 — XGBoost")
body(
    "A machine-learning model that looks at patterns like “what were sales yesterday / last week / "
    "last month,” “is today a weekend,” and a few business signals (discounts, marketing spend, "
    "and how a product's category is trending), and learns which combinations predict tomorrow's sales "
    "best. It's the most flexible of the three and, in testing on the bundled dataset, came out on top."
)
h2("How they're tested fairly")
body(
    "Every model is trained only on the older portion of the data and tested against the most recent "
    "two weeks it never saw during training — never shuffled, so no model ever accidentally “peeks” "
    "at the future. That's the same discipline a real production forecasting system needs."
)
h2("Actual results (not illustrative — this is a real run)")
table(
    ["Model", "MAE", "RMSE", "MAPE", "R²"],
    [
        ["Baseline (Seasonal Naive)", "32.93", "41.21", "4.99%", "0.834"],
        ["SARIMA", "41.63", "48.68", "6.05%", "0.768"],
        ["XGBoost (winner)", "33.72", "36.97", "5.05%", "0.866"],
    ],
    widths=[Inches(2.4), Inches(1.1), Inches(1.1), Inches(1.1), Inches(1.1)],
)
body(
    "Lower MAE/RMSE/MAPE is better; higher R² is better. XGBoost wins on RMSE, meaning it made "
    "fewer big mistakes overall — though the baseline is unusually strong here, which is a genuine "
    "and honest finding: this dataset has very regular weekly seasonality, so a sophisticated model "
    "only has a modest edge over a simple one. That's a useful, realistic result to report rather "
    "than hide.",
    italic=True, color=TEXT_SECONDARY
)

# ================================================== 8. RUN IT YOURSELF ===
h1("8. Running It Yourself")
body("Two ways to get it running:")
numbered(" — run ", "Option A: Docker (easiest) ")
body("docker compose up --build", size=10.5, color=BRAND_DARK)
body("then open http://localhost:3000 for the dashboard.")
numbered(" — start the FastAPI backend with ", "Option B: Manual ")
body("uvicorn app.main:app --reload", size=10.5, color=BRAND_DARK)
body("and the frontend with npm run dev; full step-by-step instructions are in the repository README.")

# ======================================================== 9. CONCLUSION ==
h1("9. Conclusion")
body(
    "SalesIQ demonstrates the full journey a real sales-forecasting product has to make: getting "
    "messy real-world data into a trustworthy state, understanding what it shows, building multiple "
    "forecasting models and honestly comparing them, and putting the result in front of a user through "
    "an actual working dashboard — not just a notebook full of numbers."
)
body(
    "22 automated tests, 6 sample datasets covering both the happy path and edge cases, and a written "
    "model evaluation report back up every claim made in this document.",
    space_after=4
)
link_p2 = doc.add_paragraph()
r = link_p2.add_run("github.com/icodervivek/SalesIQ")
r.font.color.rgb = BRAND
r.bold = True

out_path = "SalesIQ_Project_Report.docx"
doc.save(out_path)
print(f"Saved {out_path}")
